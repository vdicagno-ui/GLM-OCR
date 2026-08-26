"""Scrittura delle righe aggregate su foglio Excel (.xlsx).

Il foglio ha 3 colonne, nell'ordine richiesto:
    1. Data inserimento   (data odierna di inserimento)
    2. Causale
    3. Valore             (sempre NEGATIVO)

Se il file esiste già, le nuove righe vengono ACCODATE (registro progressivo);
altrimenti il file viene creato con l'intestazione.
"""

from __future__ import annotations

from datetime import date
from pathlib import Path
from typing import Iterable, List

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .aggregate import RigaAggregata

INTESTAZIONI = ["Data inserimento", "Causale", "Valore"]
NOME_FOGLIO = "Spese"


def _crea_intestazione(ws) -> None:
    ws.append(INTESTAZIONI)
    font = Font(bold=True, color="FFFFFF")
    fill = PatternFill(start_color="305496", end_color="305496", fill_type="solid")
    for col in range(1, len(INTESTAZIONI) + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = font
        cell.fill = fill
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 40
    ws.column_dimensions["C"].width = 14
    ws.freeze_panes = "A2"


def scrivi_righe(
    percorso: str,
    righe: Iterable[RigaAggregata],
    data_inserimento: date | None = None,
) -> int:
    """Scrive/accoda le righe aggregate sul file Excel indicato.

    I valori vengono scritti come NEGATIVI (spese). Restituisce il numero di
    righe scritte.
    """
    righe = list(righe)
    if data_inserimento is None:
        data_inserimento = date.today()

    percorso_p = Path(percorso)
    if percorso_p.exists():
        wb = load_workbook(percorso_p)
        ws = wb[NOME_FOGLIO] if NOME_FOGLIO in wb.sheetnames else wb.active
        # Assicura l'intestazione se il foglio è vuoto
        if ws.max_row == 1 and all(c.value is None for c in ws[1]):
            _crea_intestazione(ws)
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = NOME_FOGLIO
        _crea_intestazione(ws)

    scritte = 0
    for r in righe:
        valore_negativo = -abs(round(float(r.totale), 2))
        ws.append([data_inserimento.strftime("%d/%m/%Y"), r.causale, valore_negativo])
        # Formato valuta sulla colonna Valore appena aggiunta
        cell = ws.cell(row=ws.max_row, column=3)
        cell.number_format = '#,##0.00 "€"'
        scritte += 1

    percorso_p.parent.mkdir(parents=True, exist_ok=True)
    wb.save(percorso_p)
    return scritte
