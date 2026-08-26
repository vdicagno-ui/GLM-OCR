"""Test della logica di base (senza OCR/GUI): importi, parsing, aggregazione, Excel."""

import os
import tempfile
from datetime import date

from estrattore_spese.importi import parse_importo
from estrattore_spese.parsing import estrai_voci, is_totale
from estrattore_spese.aggregate import aggrega_per_causale
from estrattore_spese.excel_writer import scrivi_righe, INTESTAZIONI
from openpyxl import load_workbook


def test_parse_importo():
    casi = {
        "12,50": 12.50,
        "1.234,56": 1234.56,
        "€ 45": 45.0,
        "45.00": 45.0,
        "1.234": 1234.0,
        "-30,00": 30.0,
        "spesa 12,50 €": 12.50,
        "100": 100.0,
        "2.500,75 EUR": 2500.75,
    }
    for testo, atteso in casi.items():
        got = parse_importo(testo)
        assert got == atteso, f"parse_importo({testo!r}) = {got}, atteso {atteso}"
    assert parse_importo("nessun numero") is None
    print("OK  test_parse_importo")


def test_is_totale():
    for t in ["totale", "Totali", "TOTALE:", "somma", "", "  ", None, "totale spese"]:
        assert is_totale(t), f"{t!r} dovrebbe essere totale"
    for t in ["spesa", "benzina", "cena fuori"]:
        assert not is_totale(t), f"{t!r} NON dovrebbe essere totale"
    print("OK  test_is_totale")


def test_estrai_voci_json():
    j = """[
      {"causale": "Benzina", "importo": "30,00"},
      {"causale": "Spesa supermercato", "importo": "45,50"},
      {"causale": "Totale", "importo": "75,50"}
    ]"""
    voci = estrai_voci(j, origine="foto1.jpg")
    assert len(voci) == 2, f"attese 2 voci, trovate {len(voci)}"
    causali = {v.causale for v in voci}
    assert "Totale" not in causali
    assert voci[0].importo == 30.0
    print("OK  test_estrai_voci_json")


def test_estrai_voci_testo():
    testo = """Benzina 30,00
Spesa 45,50
Cena 20,00
Totale 95,50
80,00
"""
    voci = estrai_voci(testo, origine="foto2.jpg")
    causali = {v.causale.lower() for v in voci}
    assert "totale" not in " ".join(causali)
    # la riga con solo "80,00" (senza causale) va ignorata
    assert len(voci) == 3, f"attese 3 voci, trovate {len(voci)}: {[v.causale for v in voci]}"
    print("OK  test_estrai_voci_testo")


def test_aggregazione():
    j = """[
      {"causale": "Benzina", "importo": "30,00"},
      {"causale": "benzina", "importo": "20,00"},
      {"causale": "Spesa", "importo": "45,50"},
      {"causale": "totale", "importo": "95,50"}
    ]"""
    voci = estrai_voci(j)
    righe = aggrega_per_causale(voci)
    per_causale = {r.causale.lower(): r for r in righe}
    assert len(righe) == 2, f"attese 2 causali, trovate {len(righe)}"
    assert per_causale["benzina"].totale == 50.0
    assert per_causale["benzina"].numero_voci == 2
    assert per_causale["spesa"].totale == 45.50
    print("OK  test_aggregazione")


def test_excel():
    j = """[
      {"causale": "Benzina", "importo": "30,00"},
      {"causale": "benzina", "importo": "20,50"},
      {"causale": "Spesa", "importo": "45,50"}
    ]"""
    righe = aggrega_per_causale(estrai_voci(j))
    with tempfile.TemporaryDirectory() as d:
        percorso = os.path.join(d, "out.xlsx")
        # Prima scrittura: crea file + intestazione
        n = scrivi_righe(percorso, righe, data_inserimento=date(2026, 8, 26))
        assert n == 2
        wb = load_workbook(percorso)
        ws = wb.active
        assert [c.value for c in ws[1]] == INTESTAZIONI
        # I valori devono essere NEGATIVI
        riga2 = [c.value for c in ws[2]]
        assert riga2[0] == "26/08/2026"
        assert riga2[2] < 0, f"valore non negativo: {riga2[2]}"
        assert abs(riga2[2]) == 50.50
        # Seconda scrittura: accoda (registro progressivo)
        n2 = scrivi_righe(percorso, righe, data_inserimento=date(2026, 8, 27))
        wb2 = load_workbook(percorso)
        ws2 = wb2.active
        assert ws2.max_row == 1 + 2 + 2, f"righe totali attese 5, trovate {ws2.max_row}"
    print("OK  test_excel")


if __name__ == "__main__":
    test_parse_importo()
    test_is_totale()
    test_estrai_voci_json()
    test_estrai_voci_testo()
    test_aggregazione()
    test_excel()
    print("\nTutti i test superati ✅")
