"""
Oscura Documenti - GUI desktop per l'anonimizzazione di file .docx
Sostituisce nomi/cognomi (con dicitura personalizzabile per ciascuno),
date di nascita, numeri di telefono e codice fiscale in tutti i file
.docx di una cartella.
"""
import os
import re
import queue
import shutil
import threading
import tkinter as tk
from tkinter import filedialog, messagebox, scrolledtext, ttk

from docx import Document
from docx.text.paragraph import Paragraph
from dateutil import parser

TAG_NOME_DEFAULT = "[soggetto_interessato]"
DATA_SOSTITUZIONE = "00.00.00"
TAG_TELEFONO = "[numero_di_telefono_oscurato]"
TAG_CF = "[codice_fiscale_oscurato]"

MESI_IT = {
    1: "gennaio", 2: "febbraio", 3: "marzo", 4: "aprile",
    5: "maggio", 6: "giugno", 7: "luglio", 8: "agosto",
    9: "settembre", 10: "ottobre", 11: "novembre", 12: "dicembre"
}

MESI_ABBR = {
    1: "gen", 2: "feb", 3: "mar", 4: "apr",
    5: "mag", 6: "giu", 7: "lug", 8: "ago",
    9: "set", 10: "ott", 11: "nov", 12: "dic"
}

REGEX_CF = re.compile(
    r"\b[A-Z]{6}\d{2}[A-EHLMPR-T][0-9LMNPQRSTUV]{2}[A-Z]\d{3}[A-Z]\b",
    re.IGNORECASE
)

# Numeri di telefono italiani, con o senza prefisso internazionale e con
# separatori opzionali (spazi, punti, trattini).
# Si richiedono ALMENO 9 cifre: le date (max 8 cifre, es. gg.mm.aaaa) non
# possono cosi' essere scambiate per numeri di telefono ed oscurate per errore.
#  - Cellulare: 3xx + 6/7 cifre  -> 9/10 cifre
#  - Fisso:     0x  + 7/9 cifre  -> 9/11 cifre
REGEX_TELEFONO = re.compile(
    r"(?<!\d)"
    r"(?:(?:\+|00)39[\s.\-]?)?"
    r"(?:"
    r"3\d{2}(?:[\s.\-]?\d){6,7}"
    r"|"
    r"0\d(?:[\s.\-]?\d){7,9}"
    r")"
    r"(?!\d)"
)


def genera_regex_nome(nome_completo):
    # La compilazione con re.IGNORECASE rende il riconoscimento indipendente
    # dal maiuscolo/minuscolo: "MARIO ROSSI", "mario rossi" e "Mario Rossi"
    # vengono tutti individuati.
    parti = [p for p in nome_completo.split() if p]
    if not parti:
        raise ValueError("Nome e cognome non possono essere vuoti.")

    esc = [re.escape(p) for p in parti]

    if len(esc) >= 2:
        esc_primo, esc_ultimo = esc[0], esc[-1]
        ini_primo = re.escape(parti[0][0])
        ini_ultimo = re.escape(parti[-1][0])

        pattern_list = [
            r"\s+".join(esc),                    # tutti i nomi in ordine (anche 3+ parole)
            r"\s+".join(reversed(esc)),          # ordine invertito
            rf"{esc_primo}\s+{esc_ultimo}",      # primo + ultimo (salta eventuali secondi nomi)
            rf"{esc_ultimo}\s+{esc_primo}",      # ultimo + primo
            rf"{ini_primo}\.\s+{esc_ultimo}",    # M. Rossi
            rf"{ini_ultimo}\.\s+{esc_primo}",    # R. Mario
        ]
        pattern_list.extend(esc)                 # ogni parola singola
        pattern_finale = r"\b(?:" + "|".join(pattern_list) + r")\b"
    else:
        pattern_finale = rf"\b{esc[0]}\b"

    return re.compile(pattern_finale, re.IGNORECASE)


def interpreta_data(input_data):
    """Interpreta una data scritta in QUALSIASI formato, inclusi i mesi
    italiani (per esteso o abbreviati) che dateutil non riconosce da solo."""
    testo = input_data.strip().lower()
    # Sostituisce i nomi dei mesi italiani con il numero corrispondente,
    # cosi' dateutil riesce ad interpretarli. I nomi per esteso vanno
    # sostituiti prima delle abbreviazioni (es. "gennaio" prima di "gen").
    for num, nome in MESI_IT.items():
        testo = re.sub(rf"\b{nome}\b", f" {num} ", testo)
    for num, ab in MESI_ABBR.items():
        testo = re.sub(rf"\b{ab}\.?\b", f" {num} ", testo)
    # Se la data inizia con un anno a 4 cifre (formato ISO aaaa-mm-gg) va
    # interpretata come anno-mese-giorno; altrimenti si assume il giorno
    # per primo, come d'uso in Italia (gg/mm/aaaa).
    if re.match(r"^\s*\d{4}[\s./\-]", testo):
        return parser.parse(testo, yearfirst=True, dayfirst=False, fuzzy=True)
    return parser.parse(testo, dayfirst=True, fuzzy=True)


def genera_regex_data(input_data):
    data_oggetto = interpreta_data(input_data)
    giorno = data_oggetto.day
    mese = data_oggetto.month
    anno_completo = str(data_oggetto.year)
    anno_corto = anno_completo[-2:]

    # Giorno/mese con zero iniziale OPZIONALE (es. "5" oppure "05").
    gg = rf"0?{giorno}"
    mm = rf"0?{mese}"
    # Anno a 4 o 2 cifre.
    anno = rf"(?:{anno_completo}|{anno_corto})"
    # Separatori numerici ammessi: spazio, punto, barra, trattino.
    sep = r"[\s./\-]"
    # Mese testuale: per esteso o abbreviato (con punto facoltativo).
    mese_full = MESI_IT[mese]
    mese_abbr = MESI_ABBR[mese]
    mese_testo = rf"(?:{mese_full}|{mese_abbr}\.?)"

    pattern_date = [
        rf"{gg}{sep}{mm}{sep}{anno}",          # gg/mm/aaaa, g-m-aa, gg.mm.aa, gg mm aaaa ...
        rf"{anno_completo}{sep}{mm}{sep}{gg}",  # aaaa-mm-gg (ISO)
        rf"{gg}\s+{mese_testo}\s+{anno}",       # 5 marzo 1985, 05 mar 85 ...
    ]
    regex = re.compile(r"\b(?:" + "|".join(pattern_date) + r")\b", re.IGNORECASE)
    etichetta = f"{str(giorno).zfill(2)} {mese_full} {anno_completo}"
    return regex, etichetta


def applica_sostituzione_sicura(paragraph, regex, tag_copertura):
    if not paragraph or paragraph.text is None:
        return False
    testo_completo = paragraph.text
    if regex.search(testo_completo):
        nuovo_testo = regex.sub(tag_copertura, testo_completo)
        if paragraph.runs:
            paragraph.runs[0].text = nuovo_testo
            for run in paragraph.runs[1:]:
                run.text = ""
        else:
            paragraph.text = nuovo_testo
        return True
    return False


def anonimizza_singolo_file(file_path, output_path, regex_nomi, regex_date, oscura_telefoni, log):
    """Legge file_path, applica le sostituzioni e scrive il risultato in
    output_path (senza mai sovrascrivere l'originale).
    Ritorna 'modificato', 'copiato' oppure 'errore'.

    regex_nomi: lista di tuple (regex, tag_sostituzione).
    regex_date: lista di regex; ogni corrispondenza viene sostituita da DATA_SOSTITUZIONE."""
    try:
        doc = Document(file_path)
        modificato = False
        tutti_i_paragrafi = []

        tutti_i_paragrafi.extend(doc.paragraphs)

        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    tutti_i_paragrafi.extend(cell.paragraphs)

        for section in doc.sections:
            for p in section.header.paragraphs:
                tutti_i_paragrafi.append(p)
            for p_element in section.footer.paragraphs:
                tutti_i_paragrafi.append(p_element)

        for elemento in doc.element.iter():
            if elemento.tag.endswith('txbxContent'):
                for p_element in elemento.iter():
                    if p_element.tag.endswith('p'):
                        tutti_i_paragrafi.append(Paragraph(p_element, doc))

        for p in tutti_i_paragrafi:
            cambiato = False

            for regex_nome, tag_nome in regex_nomi:
                if applica_sostituzione_sicura(p, regex_nome, tag_nome):
                    cambiato = True

            # Le date vengono oscurate PRIMA dei telefoni: le date
            # configurate diventano "00.00.00" (solo 6 cifre) e non possono
            # piu' essere scambiate per un numero di telefono, mentre il
            # regex del telefono (min. 9 cifre) non tocca le date residue.
            for regex_data in regex_date:
                if applica_sostituzione_sicura(p, regex_data, DATA_SOSTITUZIONE):
                    cambiato = True

            if oscura_telefoni and applica_sostituzione_sicura(p, REGEX_TELEFONO, TAG_TELEFONO):
                cambiato = True

            if applica_sostituzione_sicura(p, REGEX_CF, TAG_CF):
                cambiato = True

            if cambiato:
                modificato = True

        if modificato:
            doc.save(output_path)
            return "modificato"
        # Nessun dato sensibile trovato: copiamo comunque l'originale nella
        # cartella di output cosi' da avere l'insieme completo dei documenti.
        shutil.copy2(file_path, output_path)
        return "copiato"
    except PermissionError:
        log(
            f"❌ Permesso negato su '{os.path.basename(file_path)}' (Errore 13). "
            "Chiudi il file se e' aperto in Word e assicurati che non sia in sola lettura."
        )
        return "errore"
    except Exception as e:
        log(f"❌ Errore durante l'elaborazione di {os.path.basename(file_path)}: {e}")
        return "errore"


def elabora_cartella(cartella, nominativi, date_input, oscura_telefoni, log):
    """nominativi: lista di tuple (nome_completo, tag_sostituzione).
    date_input: lista di stringhe data in qualsiasi formato."""
    if not os.path.exists(cartella):
        log(f"Errore: la cartella '{cartella}' non esiste.")
        return

    regex_nomi = [
        (genera_regex_nome(nome), tag.strip() or TAG_NOME_DEFAULT)
        for nome, tag in nominativi
    ]

    regex_date = []
    etichette_date = []
    for data_input in date_input:
        regex_data, etichetta_data = genera_regex_data(data_input)
        regex_date.append(regex_data)
        etichette_date.append(etichetta_data)

    log(f"✅ Nominativi configurati: {len(regex_nomi)} | Date configurate: {', '.join(etichette_date)}")
    if oscura_telefoni:
        log("✅ Oscuramento numeri di telefono attivo.")

    cartella_output = os.path.join(cartella, "Documenti_Anonimizzati")
    try:
        os.makedirs(cartella_output, exist_ok=True)
    except PermissionError:
        log(
            f"❌ Impossibile creare la cartella di output '{cartella_output}' (Errore 13). "
            "Scegli una cartella su cui hai i permessi di scrittura (es. Desktop o Documenti)."
        )
        return

    files = [
        f for f in os.listdir(cartella)
        if f.endswith('.docx') and not f.startswith('~$')
    ]

    if not files:
        log(f"Nessun file .docx trovato in '{cartella}'.")
        return

    log(f"Inizio... Elaborazione di {len(files)} file nella cartella '{cartella}'...")
    log(f"I file anonimizzati verranno salvati in: {cartella_output}\n")
    file_modificati = 0
    file_errori = 0

    for file_nome in files:
        percorso_completo = os.path.join(cartella, file_nome)
        percorso_output = os.path.join(cartella_output, file_nome)
        esito = anonimizza_singolo_file(
            percorso_completo, percorso_output, regex_nomi, regex_date, oscura_telefoni, log
        )
        if esito == "modificato":
            log(f" Anonimizzato (Rilevato testo/date/telefono/CF): {file_nome}")
            file_modificati += 1
        elif esito == "copiato":
            log(f" Copiato (Nessuna corrispondenza): {file_nome}")
        else:
            file_errori += 1

    log(f"\nFine! File anonimizzati: {file_modificati}/{len(files)}"
        + (f" | File con errori: {file_errori}" if file_errori else ""))
    log(f"Trovi tutti i documenti nella cartella: {cartella_output}")


class NominativoRow:
    def __init__(self, parent, on_remove):
        self.frame = ttk.Frame(parent)
        self.nome_var = tk.StringVar()
        self.tag_var = tk.StringVar()

        ttk.Entry(self.frame, textvariable=self.nome_var, width=26).pack(side="left", padx=(0, 4))
        ttk.Entry(self.frame, textvariable=self.tag_var, width=26).pack(side="left", padx=(0, 4))
        ttk.Button(self.frame, text="✕", width=3, command=lambda: on_remove(self)).pack(side="left")

        self.frame.pack(fill="x", pady=2)

    def valori(self):
        return self.nome_var.get().strip(), self.tag_var.get().strip()

    def destroy(self):
        self.frame.destroy()


class DataRow:
    def __init__(self, parent, on_remove):
        self.frame = ttk.Frame(parent)
        self.data_var = tk.StringVar()

        ttk.Entry(self.frame, textvariable=self.data_var, width=26).pack(side="left", padx=(0, 4))
        ttk.Button(self.frame, text="✕", width=3, command=lambda: on_remove(self)).pack(side="left")

        self.frame.pack(fill="x", pady=2)

    def valore(self):
        return self.data_var.get().strip()

    def destroy(self):
        self.frame.destroy()


class OscuraDocumentiApp:
    def __init__(self, root):
        self.root = root
        self.root.title("Oscura Documenti")
        self.root.geometry("720x700")
        self.root.minsize(640, 560)

        self.log_queue = queue.Queue()
        self.worker_thread = None
        self.nominativi_rows = []
        self.date_rows = []

        self._build_ui()
        self.root.after(100, self._poll_log_queue)

    def _build_ui(self):
        pad = {"padx": 10, "pady": 6}

        top = ttk.Frame(self.root)
        top.pack(fill="x", **pad)
        top.columnconfigure(0, weight=1)

        ttk.Label(top, text="Cartella da elaborare:").grid(row=0, column=0, sticky="w")
        self.cartella_var = tk.StringVar()
        ttk.Entry(top, textvariable=self.cartella_var).grid(row=1, column=0, sticky="we")
        ttk.Button(top, text="Sfoglia...", command=self._scegli_cartella).grid(row=1, column=1, padx=(6, 0))

        # --- Nominativi ---
        nomi_frame = ttk.LabelFrame(self.root, text="Nominativi da oscurare")
        nomi_frame.pack(fill="x", **pad)

        header = ttk.Frame(nomi_frame)
        header.pack(fill="x", pady=(4, 0))
        ttk.Label(header, text="Nome e Cognome", width=26).pack(side="left", padx=(0, 4))
        ttk.Label(header, text="Dicitura di sostituzione", width=26).pack(side="left", padx=(0, 4))

        self.nominativi_container = ttk.Frame(nomi_frame)
        self.nominativi_container.pack(fill="x")

        ttk.Button(
            nomi_frame, text="+ Aggiungi nominativo", command=self._aggiungi_nominativo
        ).pack(anchor="w", pady=(4, 6))

        # --- Date ---
        date_frame = ttk.LabelFrame(
            self.root, text=f"Date di nascita da oscurare (sostituite da \"{DATA_SOSTITUZIONE}\")"
        )
        date_frame.pack(fill="x", **pad)

        self.date_container = ttk.Frame(date_frame)
        self.date_container.pack(fill="x", pady=(4, 0))

        ttk.Button(
            date_frame, text="+ Aggiungi data", command=self._aggiungi_data
        ).pack(anchor="w", pady=(4, 6))

        # --- Opzioni ---
        opzioni_frame = ttk.Frame(self.root)
        opzioni_frame.pack(fill="x", **pad)
        self.telefoni_var = tk.BooleanVar(value=True)
        ttk.Checkbutton(
            opzioni_frame, text="Oscura anche i numeri di telefono", variable=self.telefoni_var
        ).pack(anchor="w")

        self.avvia_btn = ttk.Button(self.root, text="Avvia oscuramento", command=self._avvia)
        self.avvia_btn.pack(pady=(4, 8))

        self.log_text = scrolledtext.ScrolledText(self.root, wrap="word", height=10, state="disabled")
        self.log_text.pack(fill="both", expand=True, padx=10, pady=(0, 10))

        self._aggiungi_nominativo()
        self._aggiungi_data()

    def _aggiungi_nominativo(self):
        row = NominativoRow(self.nominativi_container, self._rimuovi_nominativo)
        self.nominativi_rows.append(row)

    def _rimuovi_nominativo(self, row):
        row.destroy()
        self.nominativi_rows.remove(row)

    def _aggiungi_data(self):
        row = DataRow(self.date_container, self._rimuovi_data)
        self.date_rows.append(row)

    def _rimuovi_data(self, row):
        row.destroy()
        self.date_rows.remove(row)

    def _scegli_cartella(self):
        percorso = filedialog.askdirectory(title="Seleziona la cartella con i file .docx")
        if percorso:
            self.cartella_var.set(percorso)

    def _log(self, messaggio):
        self.log_queue.put(messaggio)

    def _poll_log_queue(self):
        try:
            while True:
                messaggio = self.log_queue.get_nowait()
                self.log_text.configure(state="normal")
                self.log_text.insert("end", messaggio + "\n")
                self.log_text.see("end")
                self.log_text.configure(state="disabled")
        except queue.Empty:
            pass
        self.root.after(100, self._poll_log_queue)

    def _avvia(self):
        cartella = self.cartella_var.get().strip()
        nominativi = [row.valori() for row in self.nominativi_rows]
        nominativi = [(nome, tag) for nome, tag in nominativi if nome]
        date_input = [row.valore() for row in self.date_rows]
        date_input = [d for d in date_input if d]
        oscura_telefoni = self.telefoni_var.get()

        if not cartella:
            messagebox.showwarning("Campo mancante", "Seleziona la cartella da elaborare.")
            return
        if not nominativi:
            messagebox.showwarning("Campo mancante", "Inserisci almeno un nominativo da oscurare.")
            return
        if not date_input:
            messagebox.showwarning("Campo mancante", "Inserisci almeno una data da oscurare.")
            return

        try:
            for data_input in date_input:
                parser.parse(data_input, dayfirst=True, fuzzy=True)
        except Exception as e:
            messagebox.showerror("Data non valida", f"Impossibile interpretare la data '{data_input}': {e}")
            return

        self.log_text.configure(state="normal")
        self.log_text.delete("1.0", "end")
        self.log_text.configure(state="disabled")

        self.avvia_btn.configure(state="disabled")
        self.worker_thread = threading.Thread(
            target=self._esegui_elaborazione,
            args=(cartella, nominativi, date_input, oscura_telefoni),
            daemon=True
        )
        self.worker_thread.start()

    def _esegui_elaborazione(self, cartella, nominativi, date_input, oscura_telefoni):
        try:
            elabora_cartella(cartella, nominativi, date_input, oscura_telefoni, self._log)
        except Exception as e:
            self._log(f"❌ Impossibile completare l'operazione: {e}")
        finally:
            self.root.after(0, lambda: self.avvia_btn.configure(state="normal"))


def main():
    root = tk.Tk()
    OscuraDocumentiApp(root)
    root.mainloop()


if __name__ == "__main__":
    main()
