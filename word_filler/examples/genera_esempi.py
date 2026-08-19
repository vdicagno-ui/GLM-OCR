"""Genera file di esempio per provare l'applicazione.

Crea un template Word con segnaposto «da compilare …» e due file guida.
Uso:  python genera_esempi.py
"""

import os
from pathlib import Path

from docx import Document

HERE = Path(__file__).parent


def make_template(path: Path):
    doc = Document()
    doc.add_heading("ATTO DI CONFERIMENTO INCARICO", level=1)
    doc.add_paragraph(
        "Il sottoscritto da compilare nome cognome, in qualità di difensore,"
    )
    doc.add_paragraph("nel procedimento iscritto al da compilare nr. RG")
    doc.add_paragraph("accetta l'incarico conferito in data da compilare data di incarico.")

    tbl = doc.add_table(rows=2, cols=2)
    tbl.style = "Table Grid"
    tbl.rows[0].cells[0].text = "Cliente"
    tbl.rows[0].cells[1].text = "da compilare nome cognome"
    tbl.rows[1].cells[0].text = "Foro"
    tbl.rows[1].cells[1].text = "da compilare foro competente"
    doc.save(str(path))
    print(f"Template creato: {path.name}")


def make_guide(path: Path, nome, rg, data, foro):
    doc = Document()
    doc.add_paragraph("SCHEDA DATI PRATICA")
    doc.add_paragraph(f"Nome cognome: {nome}")
    doc.add_paragraph(f"Nr. RG: {rg}")
    doc.add_paragraph(f"Data di incarico: {data}")
    doc.add_paragraph(f"Foro competente: {foro}")
    doc.add_paragraph("Note aggiuntive: pratica ordinaria.")
    doc.save(str(path))
    print(f"File guida creato: {path.name}")


def main():
    os.makedirs(HERE, exist_ok=True)
    make_template(HERE / "template_esempio.docx")
    make_guide(HERE / "guida_1.docx", "Mario Rossi", "1234/2026", "10/01/2026", "Milano")
    make_guide(HERE / "guida_2.docx", "Anna Bianchi", "5678/2026", "22/02/2026", "Roma")
    print("\nFatto. Apri l'app, carica il template e i due file guida.")


if __name__ == "__main__":
    main()
