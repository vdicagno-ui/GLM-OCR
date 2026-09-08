"""Test di base per la logica core (senza GUI e senza AI)."""

import os
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))

from docx import Document

from word_filler.core import extractor, guides, pipeline, template

PHRASE_PRESET = "Etichetta «da compilare …» (senza parentesi)"


def _make_template(path):
    doc = Document()
    doc.add_paragraph("Il sottoscritto da compilare nome cognome, dichiara quanto segue.")
    doc.add_paragraph("Procedimento da compilare nr. RG")
    doc.add_paragraph("Incarico conferito il da compilare data di incarico.")
    # A table cell placeholder too.
    tbl = doc.add_table(rows=1, cols=2)
    tbl.rows[0].cells[0].text = "Cliente"
    tbl.rows[0].cells[1].text = "da compilare nome cognome"
    doc.save(path)


def test_scan_and_fill():
    with tempfile.TemporaryDirectory() as d:
        tpl = os.path.join(d, "template.docx")
        out = os.path.join(d, "out.docx")
        _make_template(tpl)

        preset = PHRASE_PRESET
        fields = template.scan_placeholders(tpl, preset)
        assert "nome cognome" in fields, fields
        assert "nr. RG" in fields, fields
        assert "data di incarico" in fields, fields
        # Deduplicated even though "nome cognome" appears twice.
        assert fields.count("nome cognome") == 1

        values = {
            "nome cognome": "Mario Rossi",
            "nr. RG": "1234/2026",
            "data di incarico": "10/01/2026",
        }
        counts = template.fill_template(tpl, values, out, preset)
        assert counts["nome cognome"] == 2  # paragraph + table cell
        assert counts["nr. RG"] == 1

        result = Document(out)
        full = "\n".join(p.text for p in result.paragraphs)
        assert "Mario Rossi" in full
        assert "1234/2026" in full
        assert "10/01/2026" in full
        assert "da compilare" not in full  # all placeholders consumed
        # Table cell replaced.
        assert result.tables[0].rows[0].cells[1].text == "Mario Rossi"
    print("test_scan_and_fill OK")


def test_run_split_placeholder():
    """Placeholder split across multiple runs must still be replaced."""
    with tempfile.TemporaryDirectory() as d:
        tpl = os.path.join(d, "t.docx")
        out = os.path.join(d, "o.docx")
        doc = Document()
        p = doc.add_paragraph()
        p.add_run("da compilare ")
        p.add_run("nome ")
        p.add_run("cognome")
        doc.save(tpl)

        fields = template.scan_placeholders(tpl, PHRASE_PRESET)
        assert fields == ["nome cognome"], fields
        template.fill_template(tpl, {"nome cognome": "Anna Bianchi"}, out, PHRASE_PRESET)
        res = Document(out)
        assert "Anna Bianchi" in res.paragraphs[0].text
        assert "da compilare" not in res.paragraphs[0].text
    print("test_run_split_placeholder OK")


def test_bracket_preset():
    with tempfile.TemporaryDirectory() as d:
        tpl = os.path.join(d, "t.docx")
        out = os.path.join(d, "o.docx")
        doc = Document()
        doc.add_paragraph("Nome: {{cliente}} - Data: {{data}}")
        doc.save(tpl)
        preset = "Doppia graffa {{campo}}"
        fields = template.scan_placeholders(tpl, preset)
        assert set(fields) == {"cliente", "data"}, fields
        template.fill_template(tpl, {"cliente": "ACME", "data": "2026"}, out, preset)
        res = Document(out)
        assert "ACME" in res.paragraphs[0].text and "2026" in res.paragraphs[0].text
    print("test_bracket_preset OK")


def test_bracket_with_prefix_default():
    """Default preset: [da compilare nome e cognome] -> field 'nome e cognome'."""
    with tempfile.TemporaryDirectory() as d:
        tpl = os.path.join(d, "t.docx")
        out = os.path.join(d, "o.docx")
        doc = Document()
        doc.add_paragraph("Il sottoscritto [da compilare nome e cognome], difensore,")
        doc.add_paragraph("procedimento [da compilare nr. RG] presso [foro competente].")
        doc.save(tpl)

        preset = template.DEFAULT_PRESET  # "Parentesi quadra [campo]"
        fields = template.scan_placeholders(tpl, preset)
        # "da compilare" prefix stripped; label with/without prefix both clean.
        assert set(fields) == {"nome e cognome", "nr. RG", "foro competente"}, fields

        values = {
            "nome e cognome": "Mario Rossi",
            "nr. RG": "1234/2026",
            "foro competente": "Milano",
        }
        template.fill_template(tpl, values, out, preset)
        res = Document(out)
        full = "\n".join(p.text for p in res.paragraphs)
        assert "Mario Rossi" in full and "1234/2026" in full and "Milano" in full
        assert "[" not in full and "]" not in full  # brackets removed
        assert "da compilare" not in full
    print("test_bracket_with_prefix_default OK")


def test_norm_key_prefix_stripping():
    from word_filler.core.template import _norm_key
    assert _norm_key("da compilare nome e cognome") == "nome e cognome"
    assert _norm_key("da  inserire   nr. RG") == "nr. RG"
    assert _norm_key("nome e cognome") == "nome e cognome"
    assert _norm_key("da compilare data di incarico.") == "data di incarico"
    print("test_norm_key_prefix_stripping OK")


def test_heuristic_extract():
    text = "Nome cognome: Mario Rossi\nNr. RG: 999/2026\nAltro dato\n"
    fields = ["nome cognome", "nr. RG", "data di incarico"]
    vals = extractor.heuristic_extract(text, fields)
    assert vals["nome cognome"] == "Mario Rossi", vals
    assert vals["nr. RG"] == "999/2026", vals
    assert vals["data di incarico"] == ""  # not present
    print("test_heuristic_extract OK")


def test_heuristic_judges_and_letterhead():
    # First line is the sender's name (letterhead): it must NOT leak into the
    # judge fields, which are anchored on their labels.
    text = (
        "Avv. Vincenzo Di Cagno - Foro di Bari\n"
        "TRIBUNALE DI BARI - Sezione Fallimentare\n"
        "Giudice Delegato: Dott. Marco Esposito\n"
        "Giudice Delegante Dott.ssa Anna Ferrari\n"  # no colon
    )
    fields = ["nome e cognome", "giudice delegato", "giudice delegante"]
    vals = extractor.heuristic_extract(text, fields)
    assert vals["giudice delegato"] == "Dott. Marco Esposito", vals
    assert vals["giudice delegante"] == "Dott.ssa Anna Ferrari", vals
    # The sender name is not grabbed by any judge field.
    assert "Di Cagno" not in vals["giudice delegato"]
    assert "Di Cagno" not in vals["giudice delegante"]
    print("test_heuristic_judges_and_letterhead OK")


def test_resolve_judges_by_context():
    """Real-world layout: judges named only by title, roles given by context."""
    text = (
        "Ill.mo Giudice Onorario di Pace\n"
        "dott.ssa Maria Rosaria OOOO\n"
        "                 su delega\n"
        "della dott.ssa Maria Luisa TTTTT\n"
        "Tribunale Civile e Penale di Bari\n"
        "          Sezione Lavoro\n"
    )
    fields = ["nome e cognome", "giudice delegato", "giudice delegante"]
    # Start from empty/weak values (as a small AI model might produce).
    values = {"nome e cognome": "Avv. Tizio", "giudice delegato": "", "giudice delegante": ""}
    out = extractor.resolve_judge_fields(text, fields, values)
    assert out["giudice delegato"] == "dott.ssa Maria Rosaria OOOO", out
    assert out["giudice delegante"] == "dott.ssa Maria Luisa TTTTT", out
    # Non-judge field untouched.
    assert out["nome e cognome"] == "Avv. Tizio"
    print("test_resolve_judges_by_context OK")


def test_resolve_judges_single_no_delega():
    """Single judge, no delegation: fill delegato, leave delegante empty."""
    text = (
        "Ill.mo Giudice Onorario di Pace\n"
        "dott. Paolo Neri\n"
        "Tribunale Civile e Penale di Bari - Sezione Lavoro\n"
    )
    fields = ["giudice delegato", "giudice delegante"]
    out = extractor.resolve_judge_fields(text, fields, {f: "" for f in fields})
    assert out["giudice delegato"] == "dott. Paolo Neri", out
    assert out["giudice delegante"] == "", out
    print("test_resolve_judges_single_no_delega OK")


def test_heuristic_label_on_own_line():
    text = "Giudice Delegato\nDott. Paolo Neri\n"
    vals = extractor.heuristic_extract(text, ["giudice delegato"])
    assert vals["giudice delegato"] == "Dott. Paolo Neri", vals
    print("test_heuristic_label_on_own_line OK")


def test_json_parsing():
    assert extractor._parse_json_object('{"a": "1"}') == {"a": "1"}
    assert extractor._parse_json_object('```json\n{"a": "1"}\n```') == {"a": "1"}
    assert extractor._parse_json_object('Ecco: {"a": "1"} fine') == {"a": "1"}
    print("test_json_parsing OK")


def test_pipeline_output_path():
    cfg = {"output_dir": "", "output_suffix": "_compilato"}
    p = pipeline.build_output_path("/tmp/guida_A.pdf", cfg)
    assert p.name == "guida_A_compilato.docx", p
    cfg2 = {"output_dir": "/tmp/out", "output_suffix": "_X"}
    p2 = pipeline.build_output_path("/x/guida_B.docx", cfg2)
    assert str(p2) == "/tmp/out/guida_B_X.docx", p2
    print("test_pipeline_output_path OK")


def test_first_page_only_docx_pagebreak():
    from docx import Document as Doc
    from docx.enum.text import WD_BREAK
    with tempfile.TemporaryDirectory() as d:
        dx = os.path.join(d, "g.docx")
        doc = Doc()
        doc.add_paragraph("TRIBUNALE DI ESEMPIO")
        doc.add_paragraph("Giudice Delegato: Dott. Rossi")
        doc.add_paragraph("Giudice Delegante: Dott. Bianchi")
        p = doc.add_paragraph()
        p.add_run().add_break(WD_BREAK.PAGE)
        doc.add_paragraph("Contenuto della seconda pagina, molto lungo, da ignorare.")
        doc.save(dx)

        full = guides.read_guide_text(dx, first_page_only=False)
        assert "seconda pagina" in full
        first = guides.read_guide_text(dx, first_page_only=True)
        assert "Giudice Delegato" in first and "Giudice Delegante" in first
        assert "seconda pagina" not in first
    print("test_first_page_only_docx_pagebreak OK")


def test_first_page_only_char_cap():
    with tempfile.TemporaryDirectory() as d:
        txt = os.path.join(d, "g.txt")
        header = "Intestazione\nGiudice Delegato: Rossi\n"
        open(txt, "w", encoding="utf-8").write(header + ("x" * 10000))
        first = guides.read_guide_text(txt, first_page_only=True)
        assert "Giudice Delegato: Rossi" in first
        assert len(first) <= guides.FIRST_PAGE_CHAR_CAP
    print("test_first_page_only_char_cap OK")


def test_guide_reading():
    with tempfile.TemporaryDirectory() as d:
        # txt
        txt = os.path.join(d, "g.txt")
        open(txt, "w", encoding="utf-8").write("Nome: Test")
        assert "Test" in guides.read_guide_text(txt)
        # docx
        dx = os.path.join(d, "g.docx")
        doc = Document()
        doc.add_paragraph("Contenuto guida docx")
        doc.save(dx)
        assert "Contenuto guida docx" in guides.read_guide_text(dx)
    print("test_guide_reading OK")


def test_pdf_form_fill():
    """Fill a fillable PDF form, keeping it editable. Skips without reportlab."""
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfgen import canvas
        from pypdf import PdfReader
    except Exception:
        print("test_pdf_form_fill SKIP (reportlab/pypdf non disponibili)")
        return

    with tempfile.TemporaryDirectory() as d:
        tpl = os.path.join(d, "modulo.pdf")
        c = canvas.Canvas(tpl, pagesize=A4)
        c.setFont("Helvetica", 12)
        c.drawString(72, 760, "Nome e cognome:")
        c.drawString(72, 680, "Giudice:")
        f = c.acroForm
        f.textfield(name="nome e cognome", x=200, y=755, width=250, height=18,
                    borderStyle="inset", value="")
        # Field identified by the placeholder text held in its value.
        f.textfield(name="campo3", x=200, y=675, width=300, height=18,
                    borderStyle="inset", value="[da compilare giudice delegato]")
        c.save()

        cfg = {
            "use_ai": False, "use_custom_regex": False,
            "preset": template.DEFAULT_PRESET, "output_dir": "",
            "output_suffix": "_compilato", "flatten_pdf": False,
        }
        fields = pipeline.scan_template_fields(tpl, cfg)
        assert set(fields) == {"nome e cognome", "giudice delegato"}, fields

        values = {"nome e cognome": "Mario Rossi",
                  "giudice delegato": "dott.ssa Maria Rosaria OOOO"}
        out = pipeline.generate_output(tpl, values, os.path.join(d, "guida.pdf"), cfg)
        assert str(out).endswith(".pdf"), out

        flds = PdfReader(str(out)).get_fields()
        # Still editable (fields present) and filled with the right values.
        assert flds is not None and len(flds) == 2, flds
        got = {fo.get("/V") for fo in flds.values()}
        assert "Mario Rossi" in got and "dott.ssa Maria Rosaria OOOO" in got, got
    print("test_pdf_form_fill OK")


if __name__ == "__main__":
    test_scan_and_fill()
    test_run_split_placeholder()
    test_bracket_preset()
    test_bracket_with_prefix_default()
    test_norm_key_prefix_stripping()
    test_heuristic_extract()
    test_heuristic_judges_and_letterhead()
    test_resolve_judges_by_context()
    test_resolve_judges_single_no_delega()
    test_heuristic_label_on_own_line()
    test_json_parsing()
    test_pipeline_output_path()
    test_first_page_only_docx_pagebreak()
    test_first_page_only_char_cap()
    test_guide_reading()
    test_pdf_form_fill()
    print("\nTutti i test superati.")
